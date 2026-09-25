"""Evidence-based normalization, resume ingestion and deterministic scoring."""
import io
import re
import hashlib
from html import unescape
from bs4 import BeautifulSoup

SKILLS = ['SQL', 'Python', 'Excel', 'Communication', 'Marketing', 'SEO', 'Content Marketing', 'Google Analytics', 'Product Analytics', 'User Research', 'Product Strategy', 'A/B Testing', 'Tableau', 'Power BI', 'LangGraph', 'Streamlit', 'GenAI', 'LLMs', 'RAG', 'Project Management', 'Sales', 'CRM', 'Copywriting', 'Figma']

def plain(value):
    return BeautifulSoup(unescape(str(value or '')), 'html.parser').get_text(' ', strip=True)

def contains(text, term):
    return bool(re.search(r'(?<!\w)' + re.escape(term.lower()) + r'(?!\w)', text.lower()))

def skills_from(text):
    return [s for s in SKILLS if contains(text, s)]

def parse_resume(name, raw):
    if len(raw) > 5 * 1024 * 1024:
        raise ValueError('Resume must be smaller than 5 MB.')
    try:
        if name.lower().endswith('.pdf'):
            from pypdf import PdfReader
            text = '\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(raw)).pages)
        elif name.lower().endswith('.docx'):
            from docx import Document
            doc = Document(io.BytesIO(raw))
            text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
        elif name.lower().endswith('.txt'):
            text = raw.decode('utf-8-sig')
        else:
            raise ValueError('Use PDF, DOCX or TXT.')
    except Exception as exc:
        raise ValueError('Could not read this resume. Use a text-based PDF, DOCX or paste the text.') from exc
    if not text.strip():
        raise ValueError('No text found. Scanned PDFs need OCR; paste your resume text instead.')
    return text.strip()

def normalize(company, source, external_id, title, description, location='', url='', employment=''):
    description = plain(description)
    # Only numeric phrases explicitly tied to experience; never infer from dates.
    found = re.search(r'(\d+(?:\.\d+)?)\s*(?:[-–]\s*\d+(?:\.\d+)?)?\s*\+?\s*years?\s+(?:of\s+)?(?:\w+\s+){0,3}experience', description, re.I)
    identity = f'{source}|{external_id}'
    return dict(id=hashlib.sha256(identity.encode()).hexdigest()[:24], company=company, source=source, external_id=str(external_id), title=plain(title), description=description, location=plain(location), url=url, employment=plain(employment), skills=skills_from(description), min_years=float(found.group(1)) if found else None)

def match(job, profile):
    weights = dict(role=30, skills=30, experience=20, location=15, preferences=5)
    dimensions = {}
    def add(key, value, explanation):
        dimensions[key] = dict(weight=weights[key], fraction=value, points=round(weights[key] * (value or 0), 1), explanation=explanation)
    roles = profile.get('roles', [])
    title_tokens = set(re.findall(r'\w+', job['title'].lower()))
    role_fit = max((len(set(re.findall(r'\w+', r.lower())) & title_tokens) / max(1, len(set(re.findall(r'\w+', r.lower())))) for r in roles), default=None)
    add('role', role_fit, 'Target-role word overlap with job title.' if roles else 'Unknown: no target roles saved.')
    requested = {s.lower() for s in job['skills']}
    owned = {s.lower() for s in profile.get('skills', [])}
    overlap, missing = sorted(requested & owned), sorted(requested - owned)
    add('skills', len(overlap)/len(requested) if requested else None, f'Matched: {", ".join(overlap) or "none"}. Missing: {", ".join(missing) or "none detected"}.' if requested else 'Unknown: no supported skill terms detected in JD.')
    years, required = profile.get('years'), job.get('min_years')
    add('experience', None if years is None or required is None else (1 if years >= required else years / max(1, required)), f'Your experience: {years}; advertised minimum: {required}. Numeric extraction needs review.')
    locations = profile.get('locations', [])
    loc = job.get('location', '')
    add('location', None if not locations or not loc else float(any(contains(loc, x) for x in locations)), f'Advertised: {loc or "unknown"}. Remote can still have country restrictions.')
    prefs = profile.get('keywords', [])
    add('preferences', sum(contains(job['description'], x) for x in prefs)/len(prefs) if prefs else None, 'Preferred keywords found in JD; unlisted preferences receive no points.')
    coverage = sum(d['weight'] for d in dimensions.values() if d['fraction'] is not None)
    return dict(score=round(sum(d['points'] for d in dimensions.values())), coverage=coverage, dimensions=dimensions, missing=missing, overlap=overlap)
